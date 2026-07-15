"""Import historical Huohua / Feigua Bilibili Excel snapshots."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from saas_crawler.storage.models import ExcelSnapshot, FeiguaBilibiliSnapshot, HuohuaUpSnapshot
from saas_crawler.storage.repositories import (
    create_creator_snapshot,
    finish_ingest_run,
    get_successful_ingest_run,
    start_ingest_run,
    upsert_creator_identity,
)


def _platform_from_path(path: Path) -> str:
    name = path.name.lower()
    if "huohua" in name or "火花" in path.name:
        return "huohua"
    if "feigua" in name or "飞瓜" in path.name:
        return "feigua"
    if "merged" in name or "合并" in path.name:
        return "merged"
    return "excel"


def _pick_first(row: dict[str, object], candidates: list[str]) -> object | None:
    lower = {str(k).lower(): v for k, v in row.items()}
    for candidate in candidates:
        if candidate in row and row[candidate] not in (None, ""):
            return row[candidate]
        value = lower.get(candidate.lower())
        if value not in (None, ""):
            return value
    return None


def import_excel_asset(session: Session, asset_id: int, path: Path) -> int:
    platform = _platform_from_path(path)
    existing = get_successful_ingest_run(session, asset_id=asset_id, source_type="excel_snapshot")
    if existing:
        return existing.imported_rows
    run = start_ingest_run(
        session,
        platform=platform,
        source_type="excel_snapshot",
        source_asset_id=asset_id,
        params={"path": str(path)},
    )
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        headers = [str(cell.value or "").strip() for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        imported = 0
        for cells in ws.iter_rows(min_row=2, values_only=True):
            row = {headers[index] or f"col_{index + 1}": value for index, value in enumerate(cells)}
            creator_id = _pick_first(row, ["B站UID", "UP主MID", "UID", "Uid", "upper_mid", "mid", "达人ID", "平台达人ID"])
            nick = _pick_first(row, ["UP主昵称", "达人昵称", "昵称", "BloggerName", "nickname", "name"])
            creator_id_text = str(creator_id or nick or f"{path.stem}:{imported + 1}")
            values = {
                "ingest_run_id": run.id,
                "platform_creator_id": creator_id_text,
                "nick_name": str(nick) if nick is not None else None,
                "row_json": row,
                "source_file": str(path),
            }
            if platform == "huohua":
                session.add(HuohuaUpSnapshot(**values))
            elif platform == "feigua":
                session.add(FeiguaBilibiliSnapshot(**values))
            else:
                session.add(ExcelSnapshot(platform=platform, **values))
            upsert_creator_identity(
                session,
                platform=platform,
                platform_creator_id=creator_id_text,
                canonical_name=str(nick) if nick is not None else None,
            )
            create_creator_snapshot(
                session,
                platform=platform,
                platform_creator_id=creator_id_text,
                ingest_run_id=run.id,
                nick_name=str(nick) if nick is not None else None,
                attributes_json=row,
                source_file=str(path),
            )
            imported += 1
        wb.close()
        finish_ingest_run(session, run, status="success", total_rows=imported, imported_rows=imported)
        session.commit()
        return imported
    except Exception as exc:
        session.rollback()
        run = session.merge(run)
        finish_ingest_run(session, run, status="failed", error_message=str(exc))
        session.commit()
        raise
